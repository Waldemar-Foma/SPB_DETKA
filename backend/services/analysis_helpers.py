from __future__ import annotations

import io
import re
from collections import Counter
from pathlib import Path

MAX_FILE_BYTES = 10 * 1024 * 1024
OKPD_RE = re.compile(r"(?<!\d)(\d{2}(?:\.\d{2}){1,4})(?!\d)")
STOPWORDS = {
    "который", "которая", "которые", "для", "при", "или", "это", "как", "что", "его", "ее", "их",
    "также", "должен", "должна", "должны", "быть", "выполнение", "оказание", "поставка", "закупка",
    "товар", "товаров", "услуг", "работ", "требования", "техническое", "задание", "согласно", "необходимо",
    "срок", "место", "исполнитель", "заказчик", "рублей", "российской", "федерации",
}


def extract_upload(file_storage) -> str:
    filename = (file_storage.filename or "").strip()
    ext = Path(filename).suffix.lower()
    raw = file_storage.read(MAX_FILE_BYTES + 1)
    if len(raw) > MAX_FILE_BYTES:
        raise ValueError("Файл больше 10 МБ")
    if not raw:
        raise ValueError("Файл пуст")
    if ext == ".txt":
        for enc in ("utf-8-sig", "utf-8", "cp1251"):
            try: return raw.decode(enc)
            except UnicodeDecodeError: pass
        raise ValueError("Не удалось определить кодировку TXT")
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
        if not text.strip():
            raise ValueError("PDF не содержит текстового слоя. Для скана нужен OCR.")
        return text
    if ext == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
        return "\n".join(parts)
    raise ValueError("Поддерживаются PDF, DOCX и TXT")


def analyse_text(text: str, filename: str | None = None) -> dict:
    clean = re.sub(r"\r\n?", "\n", text or "").strip()
    if len(clean) < 20:
        raise ValueError("Слишком мало текста для анализа")
    codes = []
    for code in OKPD_RE.findall(clean):
        if code not in codes: codes.append(code)
        if len(codes) >= 12: break
    tokens = re.findall(r"[а-яёa-z0-9-]{4,}", clean.lower())
    freq = Counter(t for t in tokens if t not in STOPWORDS and not t.isdigit())
    keywords = [word for word, _ in freq.most_common(18)]
    lines = [re.sub(r"\s+", " ", x).strip() for x in clean.splitlines() if len(x.strip()) >= 6]
    title = next((x for x in lines if len(x) <= 180), lines[0] if lines else "Новая заявка")[:180]
    return {
        "filename": filename or "текст",
        "chars": len(clean),
        "okpd2_candidates": codes,
        "keywords": keywords,
        "title": title,
        "subject": clean[:12000],
    }
